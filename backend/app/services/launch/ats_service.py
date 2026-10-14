"""
ATS Resume Checker Service (Sprint 7 V1 Scope).
Provides synchronous PDF/DOCX validation (magic bytes + size), robust text extraction,
deterministic rule-based scoring ('defined-rule-based-checks' v1.0), KIE enrichment with
graceful degradation, and analysis history persistence.
"""
from __future__ import annotations
import asyncio
from datetime import datetime, timezone
from io import BytesIO
import logging
import re
from typing import Any, Dict, List, Optional, Tuple
import uuid

from docx import Document
from fastapi import HTTPException, UploadFile, status
import httpx
from pypdf import PdfReader

from app.core.config import settings
from app.core.supabase import supabase_admin
from app.schemas.launch import (
    APIError,
    ATSAnalysisHistoryResponse,
    ATSAnalysisRecord,
    ATSAnalysisResponse,
    ATSDocumentMetadata,
    ATSFinding,
    ATSFindingCategory,
    ATSFindingSeverity,
    ATSScore,
)
from app.services.event_service import event_service

logger = logging.getLogger(__name__)

# Constants
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MIN_CHARACTER_COUNT = 50

ACTION_VERBS = {
    "accelerated", "accomplished", "achieved", "acquired", "adapted", "administered",
    "advanced", "advised", "analyzed", "architected", "authored", "automated",
    "built", "championed", "collaborated", "constructed", "coordinated", "created",
    "debugged", "delivered", "deployed", "designed", "developed", "devised",
    "directed", "discovered", "documented", "drove", "engineered", "enhanced",
    "established", "executed", "expanded", "expedited", "formulated", "founded",
    "generated", "guided", "implemented", "improved", "increased", "initiated",
    "innovated", "installed", "instituted", "integrated", "invented", "investigated",
    "launched", "led", "leveraged", "managed", "maximized", "mentored", "minimized",
    "modernized", "monitored", "negotiated", "operated", "optimized", "orchestrated",
    "organized", "overhauled", "pioneered", "planned", "produced", "programmed",
    "promoted", "reduced", "refactored", "remodeled", "reorganized", "resolved",
    "restructured", "revamped", "scaled", "simplified", "solved", "spearheaded",
    "standardized", "streamlined", "strengthened", "structured", "succeeded",
    "supervised", "tested", "tracked", "trained", "transformed", "upgraded", "validated"
}


class ATSService:
    """Service to handle ATS file validation, parsing, scoring, and persistence."""

    @staticmethod
    def _validate_file_size_and_magic_bytes(file_bytes: bytes, filename: str) -> str:
        """
        Validates file size (max 5 MB) and magic bytes for PDF and DOCX formats.
        Returns file type ('pdf' or 'docx').
        Raises HTTPException 400 with APIError if invalid.
        """
        # 1. Size check
        if len(file_bytes) > MAX_FILE_SIZE:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=APIError(
                    code="FILE_TOO_LARGE",
                    message="Resume file size exceeds the maximum limit of 5MB.",
                    details={"sizeBytes": len(file_bytes), "maxSizeBytes": MAX_FILE_SIZE},
                ).model_dump(by_alias=True),
            )

        lower_name = filename.lower()

        # 2. Magic bytes check
        # PDF magic bytes: b"%PDF-"
        is_pdf_ext = lower_name.endswith(".pdf")
        is_pdf_magic = file_bytes.startswith(b"%PDF-")

        # DOCX is a ZIP container: b"PK\x03\x04"
        is_docx_ext = lower_name.endswith(".docx")
        is_docx_magic = file_bytes.startswith(b"PK\x03\x04")

        if is_pdf_ext and is_pdf_magic:
            return "pdf"
        elif is_docx_ext and is_docx_magic:
            return "docx"
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=APIError(
                    code="UNSUPPORTED_OR_CORRUPT_FILE",
                    message="Invalid or corrupted file. Only valid PDF and DOCX files with authentic magic bytes are accepted.",
                    details={"fileName": filename},
                ).model_dump(by_alias=True),
            )

    @staticmethod
    def _extract_text(file_bytes: bytes, file_type: str) -> Tuple[str, Optional[int]]:
        """
        Extracts plain text from PDF or DOCX file bytes.
        Returns (extracted_text, page_count).
        Raises HTTP 422 DOCUMENT_PARSE_ERROR on parsing failure.
        """
        try:
            if file_type == "pdf":
                reader = PdfReader(BytesIO(file_bytes))
                page_count = len(reader.pages)
                extracted_pages = []
                for p in reader.pages:
                    txt = p.extract_text()
                    if txt:
                        extracted_pages.append(txt)
                full_text = "\n".join(extracted_pages).strip()
                return full_text, page_count
            elif file_type == "docx":
                doc = Document(BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            paragraphs.append(row_text)
                full_text = "\n".join(paragraphs).strip()
                return full_text, None
            else:
                raise ValueError(f"Unknown file type: {file_type}")
        except Exception as e:
            logger.error(f"Failed to extract text from {file_type}: {e}")
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=APIError(
                    code="DOCUMENT_PARSE_ERROR",
                    message=f"Failed to parse and extract text from document: {str(e)}",
                    details={"fileType": file_type},
                ).model_dump(by_alias=True),
            )

    @staticmethod
    def _run_rule_based_checks(text: str, page_count: Optional[int]) -> Tuple[ATSScore, List[ATSFinding]]:
        """
        Executes deterministic ATS rule-based scoring methodology ('defined-rule-based-checks' v1.0).
        Scoring categories:
          1. contact_information: 15 pts
          2. structure: 20 pts
          3. content & missing_information: 35 pts
          4. readability & formatting: 30 pts
        Total possible score: 100.0 pts.
        """
        findings: List[ATSFinding] = []
        score_contact = 0.0
        score_structure = 0.0
        score_content = 0.0
        score_readability = 0.0

        lower_text = text.lower()

        # ----------------------------------------------------------------------
        # 1. Contact Information Checks (Max 15 pts)
        # ----------------------------------------------------------------------
        # Email Check (5 pts)
        email_pattern = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
        if email_pattern.search(text):
            score_contact += 5.0
        else:
            findings.append(
                ATSFinding(
                    category="contact_information",
                    severity="critical",
                    title="Missing Email Address",
                    description="No valid email address was detected in your resume header.",
                    recommendation="Include a clear, professional email address at the very top of your resume.",
                )
            )

        # Phone Check (4 pts)
        phone_pattern = re.compile(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}")
        if phone_pattern.search(text):
            score_contact += 4.0
        else:
            findings.append(
                ATSFinding(
                    category="contact_information",
                    severity="warning",
                    title="Missing Phone Number",
                    description="No reachable phone number was found.",
                    recommendation="Add a standardized phone number with area/country code so recruiters can contact you.",
                )
            )

        # Online Profiles / Links (4 pts)
        link_pattern = re.compile(r"(linkedin\.com|github\.com|gitlab\.com|portfolio|https?://[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})", re.IGNORECASE)
        if link_pattern.search(text):
            score_contact += 4.0
        else:
            findings.append(
                ATSFinding(
                    category="contact_information",
                    severity="info",
                    title="No Professional Links Found",
                    description="No LinkedIn, GitHub, or portfolio website links were found in the text.",
                    recommendation="Include hyperlinks to your LinkedIn profile or GitHub repositories to substantiate your experience.",
                )
            )

        # Location Check (2 pts)
        location_pattern = re.compile(r"\b(remote|usa|united states|india|canada|germany|uk|ca|ny|tx|san francisco|bangalore|london|berlin|singapore)\b|,\s*[A-Z]{2}\b", re.IGNORECASE)
        if location_pattern.search(text):
            score_contact += 2.0
        else:
            score_contact += 1.0  # partial credit

        # ----------------------------------------------------------------------
        # 2. Structure & Headings Checks (Max 20 pts)
        # ----------------------------------------------------------------------
        # Education Heading (5 pts)
        if re.search(r"\b(education|academic|qualifications|degree)\b", lower_text):
            score_structure += 5.0
        else:
            findings.append(
                ATSFinding(
                    category="structure",
                    severity="warning",
                    title="Missing Education Section",
                    description="An explicitly labeled 'Education' or 'Academic' section was not identified.",
                    recommendation="Add a standard 'Education' section heading detailing your degree, university, and graduation year.",
                )
            )

        # Experience / Projects Heading (5 pts)
        if re.search(r"\b(experience|work history|employment|projects|technical projects)\b", lower_text):
            score_structure += 5.0
        else:
            findings.append(
                ATSFinding(
                    category="structure",
                    severity="critical",
                    title="Missing Experience or Projects Section",
                    description="No clear 'Experience' or 'Projects' section heading was found.",
                    recommendation="Include standard headings like 'Work Experience' or 'Projects' to ensure ATS parsers catalog your career history.",
                )
            )

        # Skills Heading (5 pts)
        if re.search(r"\b(skills|technical skills|technologies|tools|competencies)\b", lower_text):
            score_structure += 5.0
        else:
            findings.append(
                ATSFinding(
                    category="structure",
                    severity="warning",
                    title="Missing Skills Section",
                    description="A dedicated 'Skills' or 'Technologies' section heading was not detected.",
                    recommendation="Organize your core technical competencies under a clear 'Skills' heading.",
                )
            )

        # Summary / Objective Heading (5 pts)
        if re.search(r"\b(summary|professional summary|executive summary|objective|profile|about)\b", lower_text):
            score_structure += 5.0
        else:
            score_structure += 2.0  # partial credit for general flow
            findings.append(
                ATSFinding(
                    category="structure",
                    severity="info",
                    title="Missing Professional Summary",
                    description="No 'Professional Summary' or 'Objective' section was detected.",
                    recommendation="Add a brief 2-3 sentence professional summary summarizing your technical focus and top accomplishments.",
                )
            )

        # ----------------------------------------------------------------------
        # 3. Content & Measurable Impact (Max 35 pts)
        # ----------------------------------------------------------------------
        # Action Verbs (12 pts)
        words = set(re.findall(r"\b[a-zA-Z]{3,}\b", lower_text))
        matched_verbs = words.intersection(ACTION_VERBS)
        verb_count = len(matched_verbs)
        verb_score = min(12.0, verb_count * 1.5)
        score_content += verb_score
        if verb_count < 5:
            findings.append(
                ATSFinding(
                    category="content",
                    severity="warning",
                    title="Low Action Verb Density",
                    description=f"Detected only {verb_count} distinct strong action verbs in your resume.",
                    recommendation="Start bullet points with high-impact action verbs (e.g., 'Architected', 'Spearheaded', 'Optimized') rather than passive phrasing.",
                )
            )

        # Measurable Impact / Quantifiable Metrics (13 pts)
        metric_matches = re.findall(
            r"(\b\d+%\b|\$\d+[\d,]*|\b\d+\s*(?:users|clients|ms|s|gb|tb|million|k|x|teams|projects)\b|\b\d{2,}\b)",
            text,
            re.IGNORECASE,
        )
        metric_count = len(metric_matches)
        metric_score = min(13.0, metric_count * 2.5)
        score_content += metric_score
        if metric_count < 3:
            findings.append(
                ATSFinding(
                    category="missing_information",
                    severity="warning",
                    title="Few Quantifiable Impact Metrics",
                    description=f"Found only {metric_count} quantifiable data points or percentage improvements.",
                    recommendation="Include specific numbers, scale indicators, and metrics (e.g., 'reduced latency by 40%', 'served 10,000+ daily active users').",
                )
            )

        # Outcome Clarity (10 pts)
        outcome_matches = re.findall(
            r"\b(resulting in|achieved|reduced|increased|improved|saved|delivered|cut|boosted|accelerated)\b",
            lower_text,
        )
        outcome_count = len(outcome_matches)
        outcome_score = min(10.0, outcome_count * 2.5)
        score_content += outcome_score
        if outcome_count == 0:
            findings.append(
                ATSFinding(
                    category="content",
                    severity="info",
                    title="Missing Result-Driven Statements",
                    description="Experience bullet points focus heavily on duties rather than business or technical outcomes.",
                    recommendation="Structure bullets using the 'Action + Context = Measurable Result' formula.",
                )
            )

        # ----------------------------------------------------------------------
        # 4. Readability & Formatting (Max 30 pts)
        # ----------------------------------------------------------------------
        char_count = len(text)
        # Length bounds (10 pts)
        if 800 <= char_count <= 5500:
            score_readability += 10.0
        elif 400 <= char_count < 800:
            score_readability += 6.0
            findings.append(
                ATSFinding(
                    category="readability",
                    severity="info",
                    title="Resume Is Relatively Brief",
                    description=f"Resume contains {char_count} characters. A comprehensive entry-level to mid-level resume is typically between 1,200 and 4,000 characters.",
                    recommendation="Elaborate on your project deliverables, architecture decisions, and academic depth.",
                )
            )
        elif 5500 < char_count <= 8000:
            score_readability += 6.0
            findings.append(
                ATSFinding(
                    category="readability",
                    severity="info",
                    title="Resume Length is High",
                    description=f"Resume contains {char_count} characters, which may spill over onto multiple dense pages.",
                    recommendation="Condense less critical bullet points to maintain recruiter engagement and readability.",
                )
            )
        else:
            score_readability += 3.0
            findings.append(
                ATSFinding(
                    category="readability",
                    severity="warning",
                    title="Suboptimal Text Length",
                    description=f"Extracted {char_count} characters. Document length is outside standard single or two-page resume guidelines.",
                    recommendation="Ensure your resume text is focused, comprehensive, and tailored to standard length.",
                )
            )

        # Bullet Structure (10 pts)
        bullet_matches = re.findall(r"(?:^|\n)\s*[•\-\*]\s+", text)
        if len(bullet_matches) >= 3 or text.count("\n") > 15:
            score_readability += 10.0
        else:
            score_readability += 4.0
            findings.append(
                ATSFinding(
                    category="formatting",
                    severity="warning",
                    title="Lack of Clear Bullet Points",
                    description="Dense paragraphs without bullet points make resume scanning difficult for ATS algorithms and recruiters.",
                    recommendation="Break down responsibilities into distinct bullet points starting with action verbs.",
                )
            )

        # Encoding / Absence of Garbled Characters (10 pts)
        non_ascii_count = sum(1 for c in text if ord(c) > 127)
        non_ascii_ratio = non_ascii_count / max(char_count, 1)
        if non_ascii_ratio < 0.04:
            score_readability += 10.0
        else:
            score_readability += 4.0
            findings.append(
                ATSFinding(
                    category="formatting",
                    severity="warning",
                    title="Non-Standard Symbols or Formatting Detected",
                    description="Detected an elevated proportion of special characters, symbols, or custom icons.",
                    recommendation="Replace graphical icons, decorative rating bars, and complex symbols with plain UTF-8 text.",
                )
            )

        total_score_val = min(100.0, max(0.0, round(score_contact + score_structure + score_content + score_readability, 1)))

        score = ATSScore(
            value=total_score_val,
            max_value=100.0,
            methodology="defined-rule-based-checks",
            methodology_version="1.0",
        )

        return score, findings

    async def analyze_resume(
        self, learner_id: str, file: UploadFile, request_id: str
    ) -> ATSAnalysisResponse:
        """
        Synchronously validates uploaded file, extracts text, performs deterministic scoring,
        invokes KIE for qualitative refinement with graceful degradation, and persists analysis history.
        """
        # 1. Read file bytes
        file_bytes = await file.read()
        filename = file.filename or "resume"

        # 2. Validate size and magic bytes
        file_type = self._validate_file_size_and_magic_bytes(file_bytes, filename)

        # 3. Extract text
        extracted_text, page_count = self._extract_text(file_bytes, file_type)
        char_count = len(extracted_text.strip())

        # 4. Enforce minimum character threshold
        if char_count < MIN_CHARACTER_COUNT:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=APIError(
                    code="EMPTY_EXTRACTED_TEXT",
                    message="Unable to extract sufficient text. Scanned image PDFs require OCR, which is not supported in V1.",
                    details={
                        "extractedCharacterCount": char_count,
                        "minRequired": MIN_CHARACTER_COUNT,
                    },
                    request_id=request_id,
                ).model_dump(by_alias=True),
            )

        # 5. Run deterministic rule-based checks
        score, findings = self._run_rule_based_checks(extracted_text, page_count)

        # 6. Call KIE to refine executive summary and qualitative findings
        summary = (
            f"ATS Resume Analysis scored {score.value}/100 based on 'defined-rule-based-checks' v1.0. "
            f"Identified {len(findings)} area(s) for optimization across structure, quantifiable impact metrics, "
            "action verbs, and formatting."
        )

        kie_payload = {
            "request_id": request_id,
            "learner_id": learner_id,
            "document": {
                "file_name": filename,
                "file_type": file_type,
                "page_count": page_count,
                "extracted_character_count": char_count,
            },
            "extracted_text": extracted_text,
            "rule_score": score.model_dump(by_alias=True),
            "rule_findings": [f.model_dump(by_alias=True) for f in findings],
        }

        kie_url = f"{settings.KIE_BASE_URL}/api/v1/kie/career/ats/analyze"
        async with httpx.AsyncClient() as client:
            try:
                resp = await client.post(
                    kie_url,
                    json=kie_payload,
                    headers={"X-Request-ID": request_id, "Content-Type": "application/json"},
                    timeout=30.0,
                )
                if resp.status_code == 200:
                    kie_data = resp.json()
                    if kie_data.get("summary"):
                        summary = kie_data["summary"]
                    if kie_data.get("findings") and isinstance(kie_data["findings"], list):
                        refined_findings = []
                        for f in kie_data["findings"]:
                            try:
                                refined_findings.append(ATSFinding(**f))
                            except Exception:
                                pass
                        if refined_findings:
                            findings = refined_findings
            except Exception as e:
                logger.warning(f"KIE ATS refine request failed ({e}); gracefully falling back to deterministic results.")

        # 7. Persist record in database
        analysis_id = str(uuid.uuid4())
        now_str = datetime.now(timezone.utc).isoformat()
        db_record = {
            "id": analysis_id,
            "learner_id": learner_id,
            "status": "completed",
            "file_name": filename,
            "file_type": file_type,
            "storage_key": None,
            "page_count": page_count,
            "extracted_character_count": char_count,
            "score": score.model_dump(by_alias=True),
            "summary": summary,
            "findings": [f.model_dump(by_alias=True) for f in findings],
            "created_at": now_str,
        }

        try:
            res = supabase_admin.table("launch_ats_analyses").insert(db_record).execute()
            saved = res.data[0] if res.data else db_record
        except Exception as e:
            logger.error(f"Failed to persist ATS analysis record for learner {learner_id}: {e}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=APIError(
                    code="ATS_PERSISTENCE_FAILED",
                    message="Failed to save ATS analysis record to database.",
                    details={"error": str(e)},
                    request_id=request_id,
                ).model_dump(by_alias=True),
            )

        # 8. Record event
        event_service.record_event(
            learner_id=learner_id,
            event_type="ATS_ANALYSIS_COMPLETED",
            entity_type="LAUNCH_ATS_ANALYSIS",
            entity_id=analysis_id,
            payload={"score": score.value, "fileName": filename, "fileType": file_type},
        )

        metadata = ATSDocumentMetadata(
            file_name=saved["file_name"],
            file_type=saved["file_type"],
            page_count=saved.get("page_count"),
            extracted_character_count=saved.get("extracted_character_count", char_count),
        )

        record = ATSAnalysisRecord(
            id=str(saved["id"]),
            learner_id=str(saved["learner_id"]),
            status="completed",
            document=metadata,
            score=score,
            summary=summary,
            findings=findings,
            created_at=str(saved.get("created_at")),
        )

        return ATSAnalysisResponse(analysis=record)

    def get_analysis(self, learner_id: str, analysis_id: str) -> ATSAnalysisResponse:
        """Retrieves a specific ATS analysis record with strict learner ownership verification."""
        try:
            res = (
                supabase_admin.table("launch_ats_analyses")
                .select("*")
                .eq("id", analysis_id)
                .execute()
            )
            if not res.data:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=APIError(
                        code="ANALYSIS_NOT_FOUND",
                        message=f"ATS Analysis '{analysis_id}' not found.",
                    ).model_dump(by_alias=True),
                )
            row = res.data[0]
            if str(row.get("learner_id")) != str(learner_id):
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail=APIError(
                        code="UNAUTHORIZED_ANALYSIS_ACCESS",
                        message="You do not have permission to view this ATS analysis.",
                    ).model_dump(by_alias=True),
                )

            metadata = ATSDocumentMetadata(
                file_name=row["file_name"],
                file_type=row["file_type"],
                page_count=row.get("page_count"),
                extracted_character_count=row.get("extracted_character_count", 0),
            )

            score_obj = ATSScore(**row["score"]) if row.get("score") else None
            findings_objs = [ATSFinding(**f) for f in (row.get("findings") or [])]

            record = ATSAnalysisRecord(
                id=str(row["id"]),
                learner_id=str(row["learner_id"]),
                status=row.get("status", "completed"),
                document=metadata,
                score=score_obj,
                summary=row.get("summary", ""),
                findings=findings_objs,
                created_at=str(row.get("created_at")),
            )
            return ATSAnalysisResponse(analysis=record)
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Failed to fetch ATS analysis {analysis_id}: {e}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=APIError(
                    code="ATS_FETCH_FAILED",
                    message="Error retrieving ATS analysis.",
                    details={"error": str(e)},
                ).model_dump(by_alias=True),
            )

    def list_analyses(self, learner_id: str) -> ATSAnalysisHistoryResponse:
        """Lists ATS analysis history for the authenticated learner."""
        try:
            res = (
                supabase_admin.table("launch_ats_analyses")
                .select("*")
                .eq("learner_id", learner_id)
                .order("created_at", desc=True)
                .execute()
            )
            rows = res.data if res.data else []
            items = []
            for r in rows:
                metadata = ATSDocumentMetadata(
                    file_name=r["file_name"],
                    file_type=r["file_type"],
                    page_count=r.get("page_count"),
                    extracted_character_count=r.get("extracted_character_count", 0),
                )
                score_obj = ATSScore(**r["score"]) if r.get("score") else None
                findings_objs = [ATSFinding(**f) for f in (r.get("findings") or [])]
                items.append(
                    ATSAnalysisRecord(
                        id=str(r["id"]),
                        learner_id=str(r["learner_id"]),
                        status=r.get("status", "completed"),
                        document=metadata,
                        score=score_obj,
                        summary=r.get("summary", ""),
                        findings=findings_objs,
                        created_at=str(r.get("created_at")),
                    )
                )
            return ATSAnalysisHistoryResponse(items=items)
        except Exception as e:
            logger.error(f"Failed to list ATS analyses for learner {learner_id}: {e}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=APIError(
                    code="ATS_HISTORY_FETCH_FAILED",
                    message="Error retrieving ATS analysis history.",
                    details={"error": str(e)},
                ).model_dump(by_alias=True),
            )


ats_service = ATSService()
