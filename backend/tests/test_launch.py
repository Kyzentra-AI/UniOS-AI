"""
Comprehensive Test Suite for Sprint 7 Launch (AI Resume Builder & ATS Resume Checker).
Tests schemas, camelCase aliases, career context aggregation, resume services,
anti-hallucination guardrails, PDF export, ATS file validation, deterministic scoring,
KIE degradation, and API endpoints.
"""
from io import BytesIO
import json
from unittest.mock import AsyncMock, MagicMock, patch
import uuid

import docx
from fastapi import status
from fastapi.testclient import TestClient
import pypdf
import pytest
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate

from app.core.deps import get_current_user
from app.main import app
from app.schemas.launch import (
    APIError,
    ATSAnalysisRecord,
    ATSAnalysisResponse,
    ATSDocumentMetadata,
    ATSFinding,
    ATSScore,
    CreateResumeRequest,
    GenerateResumeRequest,
    ProfileResumePreviewResponse,
    ResumeAchievement,
    ResumeContact,
    ResumeDetailResponse,
    ResumeDraft,
    ResumeEducation,
    ResumeExperience,
    ResumeListResponse,
    ResumePreferences,
    ResumeProject,
    ResumeRecord,
    ResumeSource,
    ResumeStatus,
    UpdateResumeRequest,
)
from app.services.launch.ats_service import ats_service
from app.services.launch.career_context_service import career_context_service
from app.services.launch.resume_service import resume_service

client = TestClient(app)

TEST_USER_ID = "00000000-0000-0000-0000-000000000001"
TEST_LEARNER_ID = "00000000-0000-0000-0000-000000000002"
OTHER_LEARNER_ID = "00000000-0000-0000-0000-000000000003"


@pytest.fixture
def mock_get_current_user():
    prev = app.dependency_overrides.get(get_current_user)
    user = {
        "user_id": TEST_USER_ID,
        "email": "test@unios.ai",
        "role": "user",
        "payload": {},
    }
    app.dependency_overrides[get_current_user] = lambda: user
    yield
    if prev is not None:
        app.dependency_overrides[get_current_user] = prev
    else:
        app.dependency_overrides.pop(get_current_user, None)


def create_dummy_pdf(text: str = "Hello ATS World Resume") -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf)
    doc.build([Paragraph(text, getSampleStyleSheet()["Normal"])])
    return buf.getvalue()


def create_dummy_docx(text: str = "Hello ATS World Resume") -> bytes:
    doc = docx.Document()
    doc.add_heading("Resume Heading", level=1)
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==============================================================================
# 1. PYDANTIC SCHEMAS & CAMELCASE ALIAS TESTS
# ==============================================================================

def test_pydantic_camelcase_serialization():
    contact = ResumeContact(
        full_name="Alex Mercer",
        email="alex@example.com",
        phone="+1234567890",
        location="San Francisco, CA",
        linked_in_url="https://linkedin.com/in/alex",
        github_url="https://github.com/alex",
        portfolio_url="https://alex.dev",
    )
    contact_dump = contact.model_dump(by_alias=True)
    assert contact_dump["fullName"] == "Alex Mercer"
    assert contact_dump["linkedInUrl"] == "https://linkedin.com/in/alex"
    assert contact_dump["githubUrl"] == "https://github.com/alex"
    assert contact_dump["portfolioUrl"] == "https://alex.dev"

    edu = ResumeEducation(
        institution="MIT",
        degree="B.S.",
        field_of_study="Computer Science",
        start_date="2020",
        end_date="2024",
        grade="3.9",
    )
    edu_dump = edu.model_dump(by_alias=True)
    assert edu_dump["fieldOfStudy"] == "Computer Science"
    assert edu_dump["startDate"] == "2020"
    assert edu_dump["endDate"] == "2024"

    exp = ResumeExperience(
        company="Acme Corp",
        role="Software Engineer",
        is_current=True,
        description="Built distributed systems.",
    )
    exp_dump = exp.model_dump(by_alias=True)
    assert exp_dump["isCurrent"] is True

    proj = ResumeProject(
        name="UniOS AI",
        description="Career platform",
        technologies=["FastAPI", "PostgreSQL"],
        project_url="https://unios.ai",
    )
    proj_dump = proj.model_dump(by_alias=True)
    assert proj_dump["projectUrl"] == "https://unios.ai"

    draft = ResumeDraft(
        contact=contact,
        headline="AI Engineer",
        skills=["Python", "FastAPI"],
        education=[edu],
        experience=[exp],
        projects=[proj],
    )
    preview = ProfileResumePreviewResponse(
        content=draft,
        missing_fields=["contact.phone", "education[0].grade"],
    )
    preview_dump = preview.model_dump(by_alias=True)
    assert "missingFields" in preview_dump
    assert preview_dump["missingFields"] == ["contact.phone", "education[0].grade"]


def test_ats_score_validator():
    # Valid score
    score = ATSScore(value=85.5, max_value=100.0)
    assert score.value == 85.5
    assert score.max_value == 100.0

    score_dump = score.model_dump(by_alias=True)
    assert score_dump["maxValue"] == 100.0
    assert score_dump["methodologyVersion"] == "1.0"

    # Invalid score (value > max_value)
    with pytest.raises(ValueError, match="cannot exceed max_value"):
        ATSScore(value=105.0, max_value=100.0)


def test_ats_metadata_and_record_aliases():
    meta = ATSDocumentMetadata(
        file_name="resume.pdf",
        file_type="pdf",
        page_count=1,
        extracted_character_count=1500,
    )
    meta_dump = meta.model_dump(by_alias=True)
    assert meta_dump["fileName"] == "resume.pdf"
    assert meta_dump["fileType"] == "pdf"
    assert meta_dump["pageCount"] == 1
    assert meta_dump["extractedCharacterCount"] == 1500

    rec = ATSAnalysisRecord(
        id=str(uuid.uuid4()),
        learner_id=TEST_LEARNER_ID,
        document=meta,
        score=ATSScore(value=90.0, max_value=100.0),
        summary="Excellent resume",
        findings=[],
        created_at="2026-10-14T10:00:00Z",
    )
    rec_dump = rec.model_dump(by_alias=True)
    assert rec_dump["createdAt"] == "2026-10-14T10:00:00Z"
    assert rec_dump["learnerId"] == TEST_LEARNER_ID


# ==============================================================================
# 2. CAREER CONTEXT SERVICE TESTS
# ==============================================================================

@patch("app.services.launch.career_context_service.supabase_admin")
def test_career_context_preview_mapping(mock_supabase):
    # Setup mock returns
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        # 1. users
        MagicMock(data=[{"full_name": "Jane Doe", "email": "jane@example.com"}]),
        # 2. learner_profiles
        MagicMock(
            data=[
                {
                    "id": TEST_LEARNER_ID,
                    "user_id": TEST_USER_ID,
                    "primary_goal": "Full Stack Engineer",
                    "skills": ["Python", "FastAPI"],
                    "target_skills": ["Docker"],
                    "experience": "Built REST APIs for 2 years.",
                }
            ]
        ),
        # 3. academic_profiles
        MagicMock(
            data=[
                {
                    "university": "Stanford University",
                    "degree_program": "B.S. in Computer Science",
                    "academic_status": "enrolled",
                    "domain": "Software Engineering",
                    "current_year": 2022,
                    "graduation_year": 2026,
                }
            ]
        ),
    ]

    # mock build_projects query with .in_()
    proj_mock = MagicMock()
    proj_mock.data = [
        {
            "id": "proj-1",
            "title": "Cloud Dashboard",
            "description": "Real-time monitoring system",
            "skills": ["Python", "React", "PostgreSQL"],
            "status": "COMPLETED",
        }
    ]
    mock_supabase.table.return_value.select.return_value.eq.return_value.in_.return_value.order.return_value.execute.return_value = proj_mock

    # mock build_evidence query
    ev_mock = MagicMock()
    ev_mock.data = [
        {
            "id": "ev-1",
            "title": "Live Production Deployment",
            "description": "Deployed to AWS with 99.9% uptime",
            "created_at": "2026-05-01T12:00:00Z",
        }
    ]
    # mock learner_mastery query
    mast_mock = MagicMock()
    mast_mock.data = [{"topic": "FastAPI", "status": "MASTERED"}]

    def table_router(table_name):
        mock_tbl = MagicMock()
        if table_name == "build_evidence":
            mock_tbl.select.return_value.in_.return_value.order.return_value.execute.return_value = ev_mock
        elif table_name == "learner_mastery":
            mock_tbl.select.return_value.eq.return_value.in_.return_value.execute.return_value = mast_mock
        return mock_tbl

    mock_supabase.table.side_effect = lambda t: table_router(t) if t in ["build_evidence", "learner_mastery"] else MagicMock(
        select=mock_supabase.table.return_value.select
    )

    preview = career_context_service.get_profile_resume_preview(TEST_USER_ID)
    assert isinstance(preview, ProfileResumePreviewResponse)
    assert preview.content.contact.full_name == "Jane Doe"
    assert preview.content.contact.email == "jane@example.com"
    assert "B.S. in Computer Science Student" in preview.content.headline

    # Skills deduplication: Python, FastAPI, Docker, React, PostgreSQL
    skills = preview.content.skills
    assert "Python" in skills
    assert "FastAPI" in skills
    assert "React" in skills
    assert len(skills) == len(set(s.lower() for s in skills))  # Deduplicated

    # Missing fields check
    assert "contact.phone" in preview.missing_fields
    assert "contact.location" in preview.missing_fields
    assert "education[0].grade" in preview.missing_fields
    assert "experience[0].company" in preview.missing_fields
    assert "experience[0].role" in preview.missing_fields

    # Experience mapped from raw text
    assert len(preview.content.experience) == 1
    assert preview.content.experience[0].description == "Built REST APIs for 2 years."

    # Projects mapped
    assert len(preview.content.projects) == 1
    assert preview.content.projects[0].name == "Cloud Dashboard"

    # Achievements mapped
    assert len(preview.content.achievements) == 1
    assert preview.content.achievements[0].title == "Live Production Deployment"


@pytest.mark.asyncio
@patch("app.core.context_v2_assembler.ContextV2Assembler.assemble", new_callable=AsyncMock)
@patch("app.services.launch.career_context_service.supabase_admin")
async def test_assemble_career_context(mock_supabase, mock_context_assembler):
    mock_context_assembler.return_value = {"learner_id": TEST_LEARNER_ID, "learner": {}}
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        MagicMock(data=[{"id": "p1", "title": "App"}]),  # projects
        MagicMock(data=[{"topic": "Python", "status": "PROFICIENT"}]),  # mastery
    ]
    mock_supabase.table.return_value.select.return_value.in_.return_value.execute.return_value = MagicMock(
        data=[{"id": "ev-1", "title": "Cert"}]
    )

    ctx = await career_context_service.assemble_career_context(TEST_LEARNER_ID)
    assert "career_readiness" in ctx
    assert len(ctx["career_readiness"]["projects"]) == 1
    assert len(ctx["career_readiness"]["evidence"]) == 1


# ==============================================================================
# 3. RESUME SERVICE TESTS
# ==============================================================================

@patch("app.services.launch.resume_service.supabase_admin")
@patch("app.services.event_service.event_service.record_event")
def test_resume_crud(mock_record_event, mock_supabase):
    # 1. Create resume
    draft = ResumeDraft(
        contact=ResumeContact(full_name="John Developer", email="john@example.com"),
        headline="Software Engineer",
        skills=["Python"],
    )
    req = CreateResumeRequest(title="My Tech Resume", source=ResumeSource.MANUAL, content=draft)

    resume_id = str(uuid.uuid4())
    mock_row = {
        "id": resume_id,
        "learner_id": TEST_LEARNER_ID,
        "title": req.title,
        "source": "manual",
        "status": "draft",
        "content": draft.model_dump(by_alias=True),
        "created_at": "2026-10-14T12:00:00Z",
        "updated_at": "2026-10-14T12:00:00Z",
    }
    mock_supabase.table.return_value.insert.return_value.execute.return_value = MagicMock(data=[mock_row])

    resp = resume_service.create_resume(TEST_LEARNER_ID, req)
    assert resp.resume.title == "My Tech Resume"
    assert resp.resume.status == ResumeStatus.DRAFT
    mock_record_event.assert_called_once()

    # 2. Get resume (valid owner)
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(data=[mock_row])
    detail = resume_service.get_resume(TEST_LEARNER_ID, resume_id)
    assert detail.resume.id == resume_id

    # 3. Get resume (unauthorized owner -> 403)
    with pytest.raises(Exception) as excinfo:
        resume_service.get_resume(OTHER_LEARNER_ID, resume_id)
    assert "403" in str(excinfo.value) or "UNAUTHORIZED" in str(excinfo.value)

    # 4. Update resume
    update_req = UpdateResumeRequest(title="Senior Tech Resume", content=draft)
    updated_row = dict(mock_row, title="Senior Tech Resume")
    mock_supabase.table.return_value.update.return_value.eq.return_value.execute.return_value = MagicMock(data=[updated_row])
    upd_resp = resume_service.update_resume(TEST_LEARNER_ID, resume_id, update_req)
    assert upd_resp.resume.title == "Senior Tech Resume"


@pytest.mark.asyncio
@patch("app.services.launch.resume_service.career_context_service.assemble_career_context", new_callable=AsyncMock)
@patch("app.services.launch.resume_service.supabase_admin")
@patch("httpx.AsyncClient.post")
async def test_generate_resume_anti_hallucination_and_kie_success(mock_post, mock_supabase, mock_ctx):
    mock_ctx.return_value = {
        "education": {"university": "UC Berkeley"},
        "career_readiness": {"projects": [{"title": "AlphaProject"}]},
    }

    user_draft = ResumeDraft(
        contact=ResumeContact(full_name="Alice Smith", email="alice@example.com"),
        education=[ResumeEducation(institution="UC Berkeley", degree="B.S.")],
        experience=[ResumeExperience(company="Google", role="Intern", description="Wrote Go code.")],
        projects=[ResumeProject(name="AlphaProject", description="AI Tool")],
    )

    req = GenerateResumeRequest(title="Tailored Resume", source=ResumeSource.PROFILE, content=user_draft)

    # KIE response tries to hallucinate a fake school and fake company:
    kie_generated_content = {
        "contact": user_draft.contact.model_dump(),
        "education": [{"id": "edu-1", "institution": "Invented University", "degree": "M.S.", "fieldOfStudy": "CS"}],
        "experience": [{"id": "exp-1", "company": "Fictional Startup Inc", "role": "Lead", "description": "Scaled servers."}],
        "projects": [{"id": "proj-1", "name": "Invented Mystery App", "description": "Cloud tool", "technologies": ["AWS"]}],
    }

    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {"content": kie_generated_content}
    mock_post.return_value = mock_resp

    gen_id = str(uuid.uuid4())
    mock_supabase.table.return_value.insert.return_value.execute.return_value = MagicMock(
        data=[
            {
                "id": gen_id,
                "learner_id": TEST_LEARNER_ID,
                "title": req.title,
                "source": "profile",
                "status": "generated",
                "content": kie_generated_content,
                "provider": "hybrid",
                "created_at": "2026-10-14T12:00:00Z",
                "updated_at": "2026-10-14T12:00:00Z",
            }
        ]
    )

    gen_resp = await resume_service.generate_resume(TEST_LEARNER_ID, req, "req-123")
    assert gen_resp.provider == "hybrid"

    # Verify anti-hallucination guardrail cleaned the hallucinated names before db insert:
    args, kwargs = mock_supabase.table.return_value.insert.call_args
    inserted_content = args[0]["content"]
    assert inserted_content["education"][0]["institution"] == "UC Berkeley"
    assert inserted_content["experience"][0]["company"] == "Google"
    assert inserted_content["projects"][0]["name"] == "AlphaProject"


@pytest.mark.asyncio
@patch("app.services.launch.resume_service.career_context_service.assemble_career_context", new_callable=AsyncMock)
@patch("httpx.AsyncClient.post")
async def test_generate_resume_kie_unavailable_502(mock_post, mock_ctx):
    mock_ctx.return_value = {}
    user_draft = ResumeDraft(
        contact=ResumeContact(full_name="Alice Smith", email="alice@example.com"),
    )
    req = GenerateResumeRequest(title="Failing Resume", source=ResumeSource.PROFILE, content=user_draft)

    # Mock KIE raising HTTP errors on all 3 retries
    mock_post.side_effect = Exception("Connection timeout to KIE")

    with pytest.raises(Exception) as excinfo:
        await resume_service.generate_resume(TEST_LEARNER_ID, req, "req-123")

    assert "502" in str(excinfo.value) or "GENERATION_PROVIDER_UNAVAILABLE" in str(excinfo.value)


@patch("app.services.launch.resume_service.supabase_admin")
def test_export_resume_pdf_success(mock_supabase):
    draft = ResumeDraft(
        contact=ResumeContact(
            full_name="Jordan Lee",
            email="jordan@example.com",
            phone="+1-555-0199",
            location="Seattle, WA",
            linked_in_url="https://linkedin.com/in/jordan",
            github_url="https://github.com/jordan",
        ),
        headline="Senior Systems Architect",
        summary="Results-driven systems architect with 6 years experience designing cloud platforms.",
        skills=["Python", "Go", "Kubernetes", "AWS", "FastAPI"],
        experience=[
            ResumeExperience(
                company="Tech Corp",
                role="Staff Engineer",
                start_date="2022",
                end_date="Present",
                is_current=True,
                description="Architected core payment gateway handling $5M/day.\nReduced latency by 45%.",
            )
        ],
        education=[
            ResumeEducation(
                institution="University of Washington",
                degree="B.S.",
                field_of_study="Computer Science",
                start_date="2016",
                end_date="2020",
                grade="3.85 GPA",
            )
        ],
        projects=[
            ResumeProject(
                name="Distributed Cache",
                technologies=["Go", "Redis"],
                description="Built high-throughput in-memory cache system.",
                project_url="https://github.com/jordan/cache",
            )
        ],
        achievements=[
            ResumeAchievement(title="AWS Certified Solutions Architect", date="2023")
        ],
    )

    resume_id = str(uuid.uuid4())
    mock_row = {
        "id": resume_id,
        "learner_id": TEST_LEARNER_ID,
        "title": "Jordan Tech Resume",
        "source": "manual",
        "status": "generated",
        "content": draft.model_dump(by_alias=True),
        "created_at": "2026-10-14T12:00:00Z",
    }
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(data=[mock_row])

    pdf_bytes = resume_service.export_resume_pdf(TEST_LEARNER_ID, resume_id)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")

    # Read with pypdf to verify it parses correctly
    reader = pypdf.PdfReader(BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    extracted = reader.pages[0].extract_text()
    assert "Jordan Lee" in extracted
    assert "Senior Systems Architect" in extracted
    assert "Tech Corp" in extracted


# ==============================================================================
# 4. ATS RESUME CHECKER SERVICE TESTS
# ==============================================================================

def test_ats_file_validation_size_and_magic_bytes():
    # 1. File too large (> 5 MB)
    large_bytes = b"0" * (5 * 1024 * 1024 + 10)
    with pytest.raises(Exception) as exc:
        ats_service._validate_file_size_and_magic_bytes(large_bytes, "large.pdf")
    assert "FILE_TOO_LARGE" in str(exc.value)

    # 2. Corrupt / fake magic bytes (.pdf extension but text content)
    fake_pdf = b"This is just plain text masquerading as a PDF file."
    with pytest.raises(Exception) as exc:
        ats_service._validate_file_size_and_magic_bytes(fake_pdf, "fake.pdf")
    assert "UNSUPPORTED_OR_CORRUPT_FILE" in str(exc.value)

    # 3. Valid PDF magic bytes
    valid_pdf = create_dummy_pdf("Sample text")
    ft = ats_service._validate_file_size_and_magic_bytes(valid_pdf, "resume.pdf")
    assert ft == "pdf"

    # 4. Valid DOCX magic bytes
    valid_docx = create_dummy_docx("Sample text")
    ft = ats_service._validate_file_size_and_magic_bytes(valid_docx, "resume.docx")
    assert ft == "docx"


def test_ats_rule_based_scoring():
    comprehensive_resume = """
    Jane Doe
    jane.doe@example.com | +1-415-555-0199 | San Francisco, CA | https://linkedin.com/in/janedoe | https://github.com/janedoe

    PROFESSIONAL SUMMARY
    Innovative Senior Software Engineer with 5+ years of experience designing scalable distributed systems and backend architectures.

    TECHNICAL SKILLS
    Languages & Tools: Python, FastAPI, Go, Docker, Kubernetes, PostgreSQL, Redis, AWS, Terraform, CI/CD, Git

    WORK EXPERIENCE
    Lead Backend Engineer | Acme Tech | 2021 - Present
    • Spearheaded architecture modernization of microservices platform, accelerating throughput by 40%.
    • Engineered automated deployment pipeline resulting in 30% reduction in production incident rollback times.
    • Architected cloud-native payment API that processed over 10,000 daily transactions with 99.99% availability.
    • Optimized database queries and Redis caching, cutting average p99 latency from 250ms to 45ms.
    • Managed a high-performing squad of 6 engineers to deliver core enterprise features ahead of deadline.

    PROJECTS
    Distributed Task Queue | Python, Redis, Docker
    • Built distributed asynchronous task queue handling 50,000 background jobs per hour.
    • Developed fault-tolerant worker nodes resulting in zero message loss during peak load spikes.

    EDUCATION
    University of California, Berkeley
    Bachelor of Science in Computer Science | 2017 – 2021
    GPA: 3.82 / 4.0
    """

    score, findings = ats_service._run_rule_based_checks(comprehensive_resume, page_count=1)
    assert score.methodology == "defined-rule-based-checks"
    assert score.methodology_version == "1.0"
    assert score.value >= 80.0  # High score for comprehensive resume
    assert score.value <= score.max_value

    # Test sparse resume generates critical findings
    sparse_resume = "Just a short note with some random text without headers or contacts."
    sparse_score, sparse_findings = ats_service._run_rule_based_checks(sparse_resume, page_count=1)
    assert sparse_score.value < 40.0
    finding_titles = [f.title for f in sparse_findings]
    assert "Missing Email Address" in finding_titles
    assert "Missing Education Section" in finding_titles


@pytest.mark.asyncio
@patch("app.services.launch.ats_service.supabase_admin")
@patch("httpx.AsyncClient.post")
async def test_ats_analyze_resume_flow_with_pdf(mock_post, mock_supabase):
    resume_text = """
    Alex Johnson
    alex.j@example.com | +1 555 999 1234 | Seattle, WA | https://github.com/alexj
    PROFESSIONAL SUMMARY
    Results-driven software engineer specialized in cloud platforms.
    SKILLS
    Python, FastAPI, Docker, SQL, Git
    EXPERIENCE
    Software Engineer | Beta Corp | 2022 - Present
    • Developed microservices resulting in 25% improved system latency.
    • Architected data pipeline handling 2,000,000 records daily.
    EDUCATION
    University of Washington | B.S. in Computer Science | 2018 - 2022
    """
    pdf_bytes = create_dummy_pdf(resume_text)

    # Mock KIE response
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {
        "summary": "AI Refined ATS Summary: Strong technical profile.",
        "findings": [],
    }
    mock_post.return_value = mock_resp

    mock_supabase.table.return_value.insert.return_value.execute.return_value = MagicMock(
        data=[
            {
                "id": str(uuid.uuid4()),
                "learner_id": TEST_LEARNER_ID,
                "status": "completed",
                "file_name": "resume.pdf",
                "file_type": "pdf",
                "page_count": 1,
                "extracted_character_count": len(resume_text),
                "created_at": "2026-10-14T12:00:00Z",
            }
        ]
    )

    upload_file = MagicMock()
    upload_file.filename = "resume.pdf"
    upload_file.read = AsyncMock(return_value=pdf_bytes)

    analysis_resp = await ats_service.analyze_resume(TEST_LEARNER_ID, upload_file, "req-ats-1")
    assert analysis_resp.analysis.status == "completed"
    assert analysis_resp.analysis.document.file_type == "pdf"
    assert analysis_resp.analysis.summary == "AI Refined ATS Summary: Strong technical profile."


@pytest.mark.asyncio
async def test_ats_empty_text_rejection():
    # PDF with nearly no extractable text (< 50 chars)
    tiny_pdf = create_dummy_pdf("Hi")
    upload_file = MagicMock()
    upload_file.filename = "scanned.pdf"
    upload_file.read = AsyncMock(return_value=tiny_pdf)

    with pytest.raises(Exception) as exc:
        await ats_service.analyze_resume(TEST_LEARNER_ID, upload_file, "req-empty")
    assert "EMPTY_EXTRACTED_TEXT" in str(exc.value)


# ==============================================================================
# 5. FASTAPI API ENDPOINT INTEGRATION TESTS
# ==============================================================================

@patch("app.api.v1.endpoints.launch.supabase_admin")
@patch("app.api.v1.endpoints.launch.career_context_service.get_profile_resume_preview")
def test_api_profile_preview(mock_preview, mock_supabase, mock_get_current_user):
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(
        data=[{"id": TEST_LEARNER_ID}]
    )
    draft = ResumeDraft(
        contact=ResumeContact(full_name="User One", email="user1@example.com"),
    )
    mock_preview.return_value = ProfileResumePreviewResponse(content=draft, missing_fields=["contact.phone"])

    res = client.get("/api/v1/launch/resumes/profile-preview")
    assert res.status_code == 200
    data = res.json()
    assert data["content"]["contact"]["fullName"] == "User One"
    assert data["missingFields"] == ["contact.phone"]


@patch("app.api.v1.endpoints.launch.supabase_admin")
@patch("app.api.v1.endpoints.launch.resume_service.list_resumes")
def test_api_list_resumes(mock_list, mock_supabase, mock_get_current_user):
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(
        data=[{"id": TEST_LEARNER_ID}]
    )
    mock_list.return_value = ResumeListResponse(items=[])

    res = client.get("/api/v1/launch/resumes")
    assert res.status_code == 200
    assert res.json()["items"] == []


@patch("app.api.v1.endpoints.launch.supabase_admin")
@patch("app.api.v1.endpoints.launch.resume_service.create_resume")
def test_api_create_resume(mock_create, mock_supabase, mock_get_current_user):
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(
        data=[{"id": TEST_LEARNER_ID}]
    )
    draft = ResumeDraft(
        contact=ResumeContact(full_name="Candidate", email="cand@example.com"),
    )
    mock_rec = ResumeRecord(
        id=str(uuid.uuid4()),
        title="Created Resume",
        content=draft,
    )
    mock_create.return_value = ResumeDetailResponse(resume=mock_rec)

    payload = {
        "title": "Created Resume",
        "source": "manual",
        "content": {
            "contact": {"fullName": "Candidate", "email": "cand@example.com"},
        },
    }
    res = client.post("/api/v1/launch/resumes", json=payload)
    assert res.status_code == 201
    assert res.json()["resume"]["title"] == "Created Resume"


@patch("app.api.v1.endpoints.launch.supabase_admin")
@patch("app.api.v1.endpoints.launch.resume_service.export_resume_pdf")
def test_api_export_pdf_stream(mock_export, mock_supabase, mock_get_current_user):
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(
        data=[{"id": TEST_LEARNER_ID}]
    )
    dummy_pdf = create_dummy_pdf("Exported Resume Content")
    mock_export.return_value = dummy_pdf

    res = client.post(f"/api/v1/launch/resumes/{uuid.uuid4()}/export")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")


@patch("app.api.v1.endpoints.launch.supabase_admin")
@patch("app.api.v1.endpoints.launch.ats_service.analyze_resume", new_callable=AsyncMock)
def test_api_ats_upload(mock_analyze, mock_supabase, mock_get_current_user):
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(
        data=[{"id": TEST_LEARNER_ID}]
    )

    meta = ATSDocumentMetadata(file_name="test.pdf", file_type="pdf", extracted_character_count=500)
    record = ATSAnalysisRecord(
        id=str(uuid.uuid4()),
        document=meta,
        score=ATSScore(value=88.0, max_value=100.0),
        summary="ATS passed",
        findings=[],
    )
    mock_analyze.return_value = ATSAnalysisResponse(analysis=record)

    dummy_pdf = create_dummy_pdf("Test PDF Content for ATS Analysis with ample characters.")
    files = {"file": ("test.pdf", dummy_pdf, "application/pdf")}

    res = client.post("/api/v1/launch/ats/analyses", files=files, headers={"X-Request-ID": "test-req-id"})
    assert res.status_code == 200
    data = res.json()
    assert data["analysis"]["summary"] == "ATS passed"
    assert data["analysis"]["score"]["value"] == 88.0


@patch("app.api.v1.endpoints.launch.supabase_admin")
def test_api_standardized_error_format(mock_supabase, mock_get_current_user):
    # Missing learner profile -> triggers 404 with standardized APIError shape
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(data=[])

    res = client.get("/api/v1/launch/resumes", headers={"X-Request-ID": "err-req-123"})
    assert res.status_code == 404
    data = res.json()
    assert "code" in data
    assert "message" in data
    assert data["code"] == "LEARNER_PROFILE_NOT_FOUND"
    assert data.get("requestId") == "err-req-123"


def test_api_generate_resume_invalid_url_validation_error(mock_get_current_user):
    """
    Verifies that an invalid URL in the contact portfolioUrl triggers a 422 Unprocessable Entity
    with a JSON-safe response rather than crashing with a 500 TypeError on ValueError serialization.
    """
    payload = {
        "title": "My Resume",
        "source": "profile",
        "content": {
            "contact": {
                "fullName": "Test User",
                "email": "test@example.com",
                "portfolioUrl": "https://portdfolio.",
            }
        },
    }

    res = client.post("/api/v1/launch/resumes/generate", json=payload, headers={"X-Request-ID": "req-val-123"})
    assert res.status_code == 422
    data = res.json()
    assert data["success"] is False
    assert data["code"] == "VALIDATION_ERROR"
    assert "Request validation failed" in data["message"]
    assert "errors" in data
    assert len(data["errors"]) > 0

    err0 = data["errors"][0]
    assert "portfolioUrl" in err0["loc"]
    assert "Invalid URL format" in err0["msg"]
    # Check that ctx error was converted to readable string and not an un-serializable object
    assert "ctx" in err0
    assert isinstance(err0["ctx"]["error"], str)
    assert "Invalid URL format" in err0["ctx"]["error"]


@patch("app.api.v1.endpoints.launch.supabase_admin")
@patch("app.api.v1.endpoints.launch.resume_service.generate_resume", new_callable=AsyncMock)
def test_api_generate_resume_success(mock_gen, mock_supabase, mock_get_current_user):
    mock_supabase.table.return_value.select.return_value.eq.return_value.execute.return_value = MagicMock(
        data=[{"id": TEST_LEARNER_ID}]
    )

    from app.schemas.launch import GenerateResumeResponse, ResumeRecord
    contact = ResumeContact(fullName="Valid User", email="user@example.com", portfolioUrl="https://example.com/portfolio")
    draft = ResumeDraft(contact=contact, headline="Engineer")
    record = ResumeRecord(
        id=str(uuid.uuid4()),
        learner_id=TEST_LEARNER_ID,
        title="Software Engineer Resume",
        source=ResumeSource.PROFILE,
        status=ResumeStatus.GENERATED,
        content=draft,
        provider="hybrid",
        created_at="2026-10-14T12:00:00Z",
        updated_at="2026-10-14T12:00:00Z",
    )
    mock_gen.return_value = GenerateResumeResponse(resume=record, provider="hybrid")

    payload = {
        "title": "Software Engineer Resume",
        "source": "profile",
        "content": {
            "contact": {
                "fullName": "Valid User",
                "email": "user@example.com",
                "portfolioUrl": "https://example.com/portfolio",
            },
            "headline": "Engineer",
        },
    }

    res = client.post("/api/v1/launch/resumes/generate", json=payload, headers={"X-Request-ID": "gen-req-123"})
    assert res.status_code == 200
    data = res.json()
    assert "resume" in data
    assert data["resume"]["title"] == "Software Engineer Resume"
    assert data["resume"]["content"]["contact"]["portfolioUrl"] == "https://example.com/portfolio"
    assert data["provider"] == "hybrid"

